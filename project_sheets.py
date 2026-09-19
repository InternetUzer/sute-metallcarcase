"""Validated XLSX imports with server-side previews and atomic confirmation."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import secrets
import time
from zipfile import ZipFile

from flask import abort, flash, g, redirect, request, send_file, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

HEADERS = ('Работа', 'Ед. изм.', 'Плановый объём', 'Выполненный объём',
           'Цена за единицу, руб.', 'Начало работ', 'Срок завершения', 'Дата выставления счёта')
MAX_ROWS = 500


def number(value):
    if value is None or isinstance(value, bool):
        raise ValueError('Заполните объёмы и цену числами; для нулевого значения укажите 0.')
    try:
        result = Decimal(str(value).replace(' ', '').replace('\u00a0', '').replace(',', '.'))
    except InvalidOperation:
        raise ValueError('Объёмы и цена должны быть числами.') from None
    if not result.is_finite() or result < 0 or result > Decimal('1000000000000'):
        raise ValueError('Объёмы и цена должны быть от 0 до 1 000 000 000 000.')
    if result.as_tuple().exponent < -4:
        raise ValueError('Допускается не более четырёх знаков после запятой.')
    return format(result, 'f')


def day(value):
    if value is None or value == '':
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    for pattern in ('%d.%m.%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(value).strip(), pattern).date().isoformat()
        except ValueError:
            pass
    raise ValueError('Дата должна быть датой Excel либо текстом ДД.ММ.ГГГГ.')


def parse_workbook(payload):
    if not payload or len(payload) > 5 * 1024 * 1024:
        raise ValueError('Загрузите файл XLSX размером до 5 МБ.')
    try:
        with ZipFile(BytesIO(payload)) as archive:
            if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError('Файл слишком сложный. Скопируйте данные в чистый шаблон.')
        book = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception:
        raise ValueError('Не удалось прочитать XLSX. Сохраните файл заново из шаблона.') from None
    try:
        if 'Работы' not in book.sheetnames:
            raise ValueError('В файле должен быть лист «Работы» из шаблона.')
        sheet = book['Работы']
        # Ignore unreliable worksheet dimensions; inspect actual rows with a hard bound.
        sheet.reset_dimensions()
        rows = sheet.iter_rows()
        first = next(rows, ())
        if tuple(c.value for c in first) != HEADERS:
            raise ValueError('Заголовки листа «Работы» не совпадают с шаблоном.')
        output = []
        for index, cells in enumerate(rows, 2):
            if index > MAX_ROWS + 1:
                raise ValueError('В файле допускается не более 500 строк после заголовка.')
            if all(c.value is None for c in cells):
                continue
            try:
                if any(c.data_type in ('f', 'e') for c in cells):
                    raise ValueError('Формулы и ошибки Excel не принимаются; вставьте значения.')
                if len(cells) > len(HEADERS) and any(c.value is not None for c in cells[len(HEADERS):]):
                    raise ValueError('Удалите лишние столбцы с данными.')
                values = [c.value for c in cells[:len(HEADERS)]]
                values += [None] * (len(HEADERS) - len(values))
                title, unit = values[:2]
                if not isinstance(title, str) or not title.strip() or len(title) > 200:
                    raise ValueError('Укажите название работы длиной до 200 символов.')
                if not isinstance(unit, str) or not unit.strip() or len(unit) > 30:
                    raise ValueError('Укажите единицу измерения длиной до 30 символов.')
                row = [title.strip(), unit.strip(), *(number(v) for v in values[2:5]), *(day(v) for v in values[5:])]
                if row[5] and row[6] and row[5] > row[6]:
                    raise ValueError('Срок завершения раньше начала работ.')
                output.append(row)
            except ValueError as error:
                raise ValueError(f'Строка {index}: {error}') from None
        if not output:
            raise ValueError('Лист «Работы» пуст. Прежняя таблица сохранена.')
        return output
    except ValueError:
        raise
    except Exception:
        raise ValueError('Повреждённый XLSX. Сохраните данные в новом шаблоне.') from None
    finally:
        book.close()


def template_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Работы'
    sheet.append(HEADERS)
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = 'A1:H501'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243746')
    for column, width in zip('ABCDEFGH', (40, 14, 22, 22, 25, 20, 20, 28)):
        sheet.column_dimensions[column].width = width
    notes = book.create_sheet('Инструкция')
    for line in ('Заполняйте лист «Работы», не меняя заголовки. Одна строка — одна работа.',
                 'Объёмы и цена обязательны, от 0, до четырёх знаков после запятой. Валюта — рубли.',
                 'Даты необязательны: используйте ДД.ММ.ГГГГ или даты Excel.',
                 'Вставляйте значения, а не формулы. Максимум 500 строк, размер до 5 МБ.',
                 'После загрузки проверьте предпросмотр. Подтверждение заменит всю таблицу объекта.',
                 'Неверный или пустой файл не изменит существующую таблицу.'):
        notes.append([line])
    notes.column_dimensions['A'].width = 115
    result = BytesIO()
    book.save(result)
    result.seek(0)
    return result


def register_sheets(app, db, protected, get_project, render):
    with app.app_context():
        db().executescript('''
            CREATE TABLE IF NOT EXISTS project_sheets(
                project_id INTEGER PRIMARY KEY REFERENCES projects(id),
                revision INTEGER NOT NULL, payload TEXT NOT NULL,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
            CREATE TABLE IF NOT EXISTS sheet_previews(
                token TEXT PRIMARY KEY, project_id INTEGER NOT NULL REFERENCES projects(id),
                admin_id INTEGER NOT NULL REFERENCES users(id), base_revision INTEGER NOT NULL,
                payload TEXT NOT NULL, expires INTEGER NOT NULL);
        ''')
        db().commit()

    @app.get('/admin/excel-template.xlsx')
    @protected(admin=True)
    def excel_template():
        return send_file(template_bytes(), as_attachment=True, download_name='project-template.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.route('/admin/projects/<int:project_id>/excel', methods=['GET', 'POST'])
    @protected(admin=True)
    def excel_import(project_id):
        project = get_project(project_id)
        error = None
        if request.method == 'POST':
            try:
                upload = request.files.get('workbook')
                if not upload or not upload.filename.lower().endswith('.xlsx'):
                    raise ValueError('Выберите файл XLSX, заполненный по шаблону.')
                rows = parse_workbook(upload.read(5 * 1024 * 1024 + 1))
                token = secrets.token_urlsafe(32)
                current = db().execute('SELECT revision FROM project_sheets WHERE project_id=?', (project_id,)).fetchone()
                db().execute('DELETE FROM sheet_previews WHERE expires < ? OR (project_id=? AND admin_id=?)',
                             (int(time.time()), project_id, g.user['id']))
                db().execute('INSERT INTO sheet_previews VALUES(?,?,?,?,?,?)',
                             (token, project_id, g.user['id'], current['revision'] if current else 0,
                              json.dumps(rows, ensure_ascii=False), int(time.time()) + 1800))
                db().commit()
                return redirect(url_for('excel_preview', project_id=project_id, token=token))
            except ValueError as exc:
                db().rollback()
                error = str(exc)
        return render('excel_import.html', 'Импорт Excel — Металл-Каркас', project=project, error=error), 422 if error else 200

    @app.route('/admin/projects/<int:project_id>/excel/<token>', methods=['GET', 'POST'])
    @protected(admin=True)
    def excel_preview(project_id, token):
        project = get_project(project_id)
        connection = db()
        if request.method == 'POST':
            connection.execute('BEGIN IMMEDIATE')
        preview = connection.execute('SELECT * FROM sheet_previews WHERE token=? AND project_id=? AND admin_id=? AND expires>?',
                                     (token, project_id, g.user['id'], int(time.time()))).fetchone()
        if not preview:
            connection.rollback()
            abort(410)
        if request.method == 'POST':
            if request.form.get('action') == 'cancel':
                connection.execute('DELETE FROM sheet_previews WHERE token=?', (token,))
                connection.commit()
                flash('Импорт отменён. Прежняя таблица сохранена.')
                return redirect(url_for('project', project_id=project_id))
            if request.form.get('action') != 'confirm':
                connection.rollback()
                abort(400)
            current = connection.execute('SELECT revision FROM project_sheets WHERE project_id=?', (project_id,)).fetchone()
            if (current['revision'] if current else 0) != preview['base_revision']:
                connection.rollback()
                return render('excel_import.html', 'Таблица уже обновлена', project=project,
                              error='Другой импорт уже изменил таблицу. Загрузите файл заново и проверьте новый предпросмотр.'), 409
            try:
                connection.execute('''INSERT INTO project_sheets(project_id,revision,payload) VALUES(?,?,?)
                    ON CONFLICT(project_id) DO UPDATE SET revision=excluded.revision,
                    payload=excluded.payload,updated_at=CURRENT_TIMESTAMP''',
                    (project_id, preview['base_revision'] + 1, preview['payload']))
                connection.execute('DELETE FROM sheet_previews WHERE token=?', (token,))
                connection.commit()
            except Exception:
                connection.rollback()
                raise
            flash('Таблица обновлена и доступна клиенту.')
            return redirect(url_for('project', project_id=project_id))
        return render('excel_preview.html', 'Предпросмотр Excel — Металл-Каркас', project=project,
                      sheet_rows=json.loads(preview['payload']), sheet_headers=HEADERS)
