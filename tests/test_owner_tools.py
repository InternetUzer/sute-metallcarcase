import io
import sqlite3
import tarfile
import time

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image
from project_sheets import HEADERS, parse_workbook
from test_site import app, client, connection, csrf, login_as, seed


def workbook(rows=None):
    book = Workbook()
    book.active.title = 'Работы'
    book.active.append(HEADERS)
    for row in rows if rows is not None else [['Монтаж', 'т', 10, 2.5, 1200, '01.10.2026', '31.10.2026', '05.10.2026']]:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    return output


def upload(client, rows=None):
    return client.post('/admin/projects/1/excel', data={
        'csrf_token': csrf(client, '/admin/projects/1/excel'),
        'workbook': (workbook(rows), 'work.xlsx')})


def confirm(client, path, action='confirm'):
    return client.post(path, data={'csrf_token': csrf(client, path), 'action': action})


def current(app):
    with connection(app) as db:
        row = db.execute('SELECT * FROM project_sheets WHERE project_id=1').fetchone()
        return dict(row) if row else None


def test_preview_confirm_and_client_isolation(app, client):
    seed(app); login_as(client, 1)
    template = client.get('/admin/excel-template.xlsx')
    assert tuple(c.value for c in load_workbook(io.BytesIO(template.data))['Работы'][1]) == HEADERS
    preview = upload(client)
    assert preview.status_code == 302
    assert current(app) is None
    assert '1200' in client.get(preview.location).text
    assert confirm(client, preview.location).status_code == 302
    assert current(app)['revision'] == 1
    assert confirm(client, preview.location).status_code == 410
    login_as(client, 2)
    assert '1200' in client.get('/cabinet/projects/1').text
    assert client.get('/cabinet/projects/2').status_code == 404
    assert client.get('/admin/excel-template.xlsx').status_code == 403
    assert client.get('/admin/projects/1/excel').status_code == 403
    assert confirm(client, preview.location).status_code == 403
    login_as(client, 3)
    assert client.get('/cabinet/projects/1').status_code == 404


@pytest.mark.parametrize('rows', [[], [['Монтаж', 'т', -1, 0, 100]],
    [['Монтаж', 'т', 1, 0, '=1+1']], [['Монтаж', 'т', 1, 0, 'NaN']],
    [['Монтаж', 'т', 1, 0, 100, 'не дата']],
    [['Монтаж', 'т', 1, 0, 100, '02.10.2026', '01.10.2026']],
    [['Монтаж', 'т', 1, 0, 100, None, None, None, 'лишнее']]])
def test_invalid_import_preserves_previous_table(app, client, rows):
    seed(app); login_as(client, 1)
    confirm(client, upload(client).location)
    previous = current(app)
    assert upload(client, rows).status_code == 422
    assert current(app) == previous


def test_cancel_expiry_wrong_project_and_concurrent_import(app, client):
    seed(app); login_as(client, 1)
    preview = upload(client).location
    assert confirm(client, preview.replace('/projects/1/', '/projects/2/')).status_code == 410
    assert confirm(client, preview, 'cancel').status_code == 302
    assert current(app) is None
    preview = upload(client).location
    with connection(app) as db:
        db.execute('UPDATE sheet_previews SET expires=?', (int(time.time()) - 1,))
    assert confirm(client, preview).status_code == 410
    preview = upload(client).location
    with connection(app) as db:
        db.execute("INSERT INTO project_sheets(project_id,revision,payload) VALUES(1,1,'[]')")
    assert confirm(client, preview).status_code == 409
    assert current(app)['payload'] == '[]'


def test_preview_bound_to_owner_and_atomic_rollback(app, client):
    seed(app); login_as(client, 1)
    confirm(client, upload(client).location)
    previous = current(app)
    preview = upload(client, [['Другие работы', 'м', 2, 1, 300]]).location
    with connection(app) as db:
        db.execute("INSERT INTO users(id,email,name,password,role) VALUES(4,'other@owner.ru','Other','unused','admin')")
        db.execute("CREATE TRIGGER fail_confirm BEFORE DELETE ON sheet_previews BEGIN SELECT RAISE(ABORT, 'simulated failure'); END")
    login_as(client, 4)
    assert client.get(preview).status_code == 410
    login_as(client, 1)
    with pytest.raises(sqlite3.IntegrityError):
        confirm(client, preview)
    assert current(app) == previous


def test_corrupt_empty_and_oversized_workbooks():
    for payload in (b'not a workbook', b'', b'x' * (5 * 1024 * 1024 + 1)):
        with pytest.raises(ValueError):
            parse_workbook(payload)
    with pytest.raises(ValueError, match='500'):
        parse_workbook(workbook([['Работа', 'шт', 1, 0, 100]] * 501).getvalue())
    assert parse_workbook(workbook([['Работа', 'шт', '1,25', 0, '100,50']]).getvalue())[0][2:5] == ['1.25', '0', '100.50']


def photo_bytes():
    result = io.BytesIO()
    Image.new('RGB', (40, 30), 'red').save(result, 'JPEG')
    result.seek(0)
    return result


def test_photo_upload_publication_and_backup(app, client, tmp_path):
    seed(app); login_as(client, 1)
    token = csrf(client, '/admin/photos')
    assert client.post('/admin/photos', data={'csrf_token': token, 'title': 'Монтаж каркаса',
        'photo': (photo_bytes(), '../../photo.jpg')}).status_code == 302
    with connection(app) as db:
        photo = db.execute('SELECT * FROM photos').fetchone()
    url = '/media/' + photo['storage_name']
    assert client.get(url).status_code == 404
    preview = client.get('/admin/photos/1/image')
    assert preview.mimetype == 'image/webp'
    assert Image.open(io.BytesIO(preview.data)).size == (40, 30)
    assert photo['storage_name'] in client.get('/admin').text
    assert client.post('/admin', data={'csrf_token': token, 'action': 'showcase',
        'image': url, 'title': 'Новый объект', 'description': 'Монтаж'}).status_code == 302
    anonymous = app.test_client()
    assert anonymous.get(url).status_code == 200
    assert url in anonymous.get('/produkciya').text
    assert anonymous.get('/admin/photos/1/image').status_code == 302
    result = app.test_cli_runner().invoke(args=['backup', str(tmp_path / 'backup.tar.gz')])
    assert result.exit_code == 0, result.output
    with tarfile.open(tmp_path / 'backup.tar.gz') as archive:
        assert 'media/' + photo['storage_name'] in archive.getnames()
    client.post('/admin', data={'csrf_token': token, 'action': 'unpublish', 'showcase_id': 1})
    assert anonymous.get(url).status_code == 404
    login_as(client, 2)
    assert client.get('/admin/photos').status_code == 403
    assert client.get('/admin/photos/1/image').status_code == 403
    assert client.post('/admin/photos', data={'csrf_token': token}).status_code == 403


def test_photo_validation_and_csrf_preserve_storage(app, client):
    seed(app); login_as(client, 1)
    assert client.post('/admin/photos', data={'title': 'No CSRF'}).status_code == 400
    token = csrf(client, '/admin/photos')
    for payload in (b'<svg><script>alert(1)</script></svg>', b'invalid jpeg'):
        assert client.post('/admin/photos', data={'csrf_token': token, 'title': 'Фото',
            'photo': (io.BytesIO(payload), 'photo.jpg')}).status_code == 422
    with connection(app) as db:
        assert db.execute('SELECT count(*) FROM photos').fetchone()[0] == 0
